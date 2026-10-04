"""M2 摄取模块测试：标题识别、切块器、文档存储、解析器、流水线。"""

import os
import tempfile

import pytest

from knowresearch.adapters import (
    BM25FullTextStore,
    InMemoryVectorStore,
    LocalObjectStorage,
    MockEmbedding,
    SQLiteDocumentStore,
)
from knowresearch.core.schemas import Document, SectionBlock
from knowresearch.ingestion.chunker import Chunker
from knowresearch.ingestion.parsers.heading_patterns import match_heading
from knowresearch.ingestion.pipeline import IngestPipeline


# ================================================================
# 标题模式识别
# ================================================================


def test_heading_arabic_multi_level():
    assert match_heading("1.1 研究背景") == (2, "研究背景")
    assert match_heading("2.3.1 实验设计") == (3, "实验设计")
    assert match_heading("  1.2.3 方法  ") == (3, "方法")


def test_heading_arabic_single():
    assert match_heading("1 引言") == (1, "引言")
    assert match_heading("2. 相关工作") == (1, "相关工作")


def test_heading_parentheses():
    assert match_heading("(1) 研究问题") == (1, "研究问题")
    assert match_heading("（2）方法") == (1, "方法")
    assert match_heading("3) 实验") == (1, "实验")


def test_heading_chinese_number():
    assert match_heading("一、绪论") == (1, "绪论")
    assert match_heading("二、相关工作") == (1, "相关工作")


def test_heading_roman():
    assert match_heading("I Introduction") == (1, "Introduction")
    assert match_heading("II Related Work") == (1, "Related Work")


def test_heading_letter():
    assert match_heading("A 附录") == (1, "附录")
    assert match_heading("B 数据") == (1, "数据")


def test_heading_not_matching():
    assert match_heading("这是一段正文，不是标题") is None
    assert match_heading("") is None
    assert match_heading("12345") is None


# ================================================================
# 切块器
# ================================================================


def test_chunker_basic():
    section = SectionBlock(
        content="这是第一段。" * 60,
        metadata={"document_id": "doc1", "section_level": 1, "section_title": "测试"},
    )
    chunker = Chunker(target_size=100, overlap=20)
    chunks = chunker.chunk_section(section)
    assert len(chunks) > 1
    for c in chunks:
        assert c.metadata["document_id"] == "doc1"
        assert c.metadata["section_id"] == section.id
        assert len(c.content) > 0


def test_chunker_short_text_single_chunk():
    section = SectionBlock(
        content="短文本，不需要切分。",
        metadata={"document_id": "doc1", "section_level": 1},
    )
    chunker = Chunker(target_size=500, overlap=100)
    chunks = chunker.chunk_section(section)
    assert len(chunks) == 1
    assert chunks[0].content == "短文本，不需要切分。"


def test_chunker_empty_content():
    section = SectionBlock(
        content="",
        metadata={"document_id": "doc1", "section_level": 1},
    )
    chunker = Chunker()
    chunks = chunker.chunk_section(section)
    assert chunks == []


def test_chunker_overlap_preserved():
    text = "A" * 400 + "\n\n" + "B" * 400
    section = SectionBlock(
        content=text,
        metadata={"document_id": "doc1", "section_level": 1},
    )
    chunker = Chunker(target_size=300, overlap=50)
    chunks = chunker.chunk_section(section)
    assert len(chunks) >= 2
    joined = "".join(c.content for c in chunks)
    assert "A" in joined and "B" in joined


# ================================================================
# DocumentStore
# ================================================================


@pytest.fixture
def doc_store():
    db_path = os.path.join(tempfile.gettempdir(), "test_docstore.db")
    if os.path.exists(db_path):
        os.remove(db_path)
    store = SQLiteDocumentStore(db_path)
    yield store
    from knowresearch.adapters.sqlite_base import close_connection

    close_connection(db_path)
    if os.path.exists(db_path):
        os.remove(db_path)


def test_document_store_crud(doc_store):
    doc = Document(
        content="",
        source="/tmp/test.pdf",
        metadata={"file_name": "test.pdf", "file_hash": "abc123", "file_type": "pdf"},
    )
    doc_store.save_document(doc)

    fetched = doc_store.get_document(doc.id)
    assert fetched is not None
    assert fetched.file_name == "test.pdf"
    assert fetched.file_hash == "abc123"

    by_hash = doc_store.get_document_by_hash("abc123")
    assert by_hash is not None
    assert by_hash.id == doc.id

    docs = doc_store.list_documents()
    assert len(docs) == 1


def test_document_store_sections(doc_store):
    doc = Document(
        content="",
        source="/tmp/test.pdf",
        metadata={"file_name": "test.pdf", "file_hash": "hash1"},
    )
    doc_store.save_document(doc)

    sections = [
        SectionBlock(
            content="引言内容",
            source="/tmp/test.pdf",
            metadata={"document_id": doc.id, "section_title": "引言", "section_level": 1, "order_index": 0},
        ),
        SectionBlock(
            content="方法内容",
            source="/tmp/test.pdf",
            metadata={"document_id": doc.id, "section_title": "方法", "section_level": 1, "order_index": 1},
        ),
    ]
    doc_store.save_sections(sections)

    fetched = doc_store.get_sections(doc.id)
    assert len(fetched) == 2
    assert fetched[0].section_title == "引言"
    assert fetched[1].section_title == "方法"


def test_document_store_delete_cascade(doc_store):
    doc = Document(
        content="",
        source="/tmp/test.pdf",
        metadata={"file_name": "test.pdf", "file_hash": "hash2"},
    )
    doc_store.save_document(doc)
    section = SectionBlock(
        content="内容",
        source="/tmp/test.pdf",
        metadata={"document_id": doc.id, "section_title": "标题", "section_level": 1, "order_index": 0},
    )
    doc_store.save_sections([section])

    doc_store.delete_document(doc.id)
    assert doc_store.get_document(doc.id) is None
    assert doc_store.get_sections(doc.id) == []


# ================================================================
# 解析器（动态生成测试文件）
# ================================================================


def _create_test_pdf(path: str) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    lines = [
        "1 Introduction",
        "This is the introduction content.",
        "1.1 Background",
        "Detailed background description.",
        "2 Method",
        "Method section content.",
        "References",
        "[1] Author. Paper Title. Journal, 2024.",
    ]
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontname="helv", fontsize=12)
        y += 20
    doc.save(path)
    doc.close()


def _create_test_docx(path: str) -> None:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    h1 = doc.add_heading("引言", level=1)
    doc.add_paragraph("这是引言部分的内容。")
    doc.add_heading("研究背景", level=2)
    doc.add_paragraph("研究背景的详细说明。")
    doc.add_heading("方法", level=1)
    doc.add_paragraph("方法部分内容。")
    doc.save(path)


def test_pdf_parser():
    from knowresearch.ingestion.parsers import PDFParser

    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = os.path.join(tmpdir, "test.pdf")
        _create_test_pdf(pdf_path)

        parser = PDFParser()
        doc, sections = parser.parse(pdf_path)

        assert doc.file_type == "pdf"
        assert doc.total_pages == 1
        assert len(sections) > 0

        titles = [s.section_title for s in sections]
        assert "Introduction" in titles
        assert "Background" in titles
        assert "Method" in titles


def test_word_parser():
    from knowresearch.ingestion.parsers import WordParser

    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "test.docx")
        _create_test_docx(docx_path)

        parser = WordParser()
        doc, sections = parser.parse(docx_path)

        assert doc.file_type == "docx"
        assert len(sections) > 0

        titles = [s.section_title for s in sections]
        assert "引言" in titles
        assert "研究背景" in titles
        assert "方法" in titles


# ================================================================
# IngestPipeline 集成测试
# ================================================================


def test_pipeline_end_to_end():
    with tempfile.TemporaryDirectory() as tmpdir:
        source_dir = os.path.join(tmpdir, "source")
        os.makedirs(source_dir)
        pdf_path = os.path.join(source_dir, "test.pdf")
        _create_test_pdf(pdf_path)

        db_path = os.path.join(tmpdir, "test.db")
        file_dir = os.path.join(tmpdir, "files")

        doc_store = SQLiteDocumentStore(db_path)
        vec_store = InMemoryVectorStore(dimension=384)
        ft_store = BM25FullTextStore()
        obj_store = LocalObjectStorage(file_dir)
        embedding = MockEmbedding()

        pipeline = IngestPipeline(
            document_store=doc_store,
            vector_store=vec_store,
            fulltext_store=ft_store,
            object_storage=obj_store,
            embedding_provider=embedding,
        )

        try:
            summary = pipeline.ingest_directory(source_dir)

            assert summary.total == 1
            assert summary.success == 1
            assert summary.results[0].section_count > 0
            assert summary.results[0].chunk_count > 0

            docs = doc_store.list_documents()
            assert len(docs) == 1

            sections = doc_store.get_sections(docs[0].id)
            assert len(sections) > 0
        finally:
            from knowresearch.adapters.sqlite_base import close_connection

            close_connection(db_path)


def test_pipeline_dedup():
    with tempfile.TemporaryDirectory() as tmpdir:
        source_dir = os.path.join(tmpdir, "source")
        os.makedirs(source_dir)
        pdf_path = os.path.join(source_dir, "test.pdf")
        _create_test_pdf(pdf_path)

        db_path = os.path.join(tmpdir, "test.db")
        file_dir = os.path.join(tmpdir, "files")

        doc_store = SQLiteDocumentStore(db_path)
        vec_store = InMemoryVectorStore(dimension=384)
        ft_store = BM25FullTextStore()
        obj_store = LocalObjectStorage(file_dir)
        embedding = MockEmbedding()

        pipeline = IngestPipeline(
            document_store=doc_store,
            vector_store=vec_store,
            fulltext_store=ft_store,
            object_storage=obj_store,
            embedding_provider=embedding,
        )

        try:
            summary1 = pipeline.ingest_directory(source_dir)
            assert summary1.success == 1

            summary2 = pipeline.ingest_directory(source_dir)
            assert summary2.skipped == 1
        finally:
            from knowresearch.adapters.sqlite_base import close_connection

            close_connection(db_path)