"""P2 真实 PostgreSQL/Redis 集成测试。

默认跳过，Docker Compose 服务启动后手动开启：
    $env:RUN_P2_INTEGRATION='1'
    python -m pytest tests/test_p2_integration.py -q
"""

import os
import shutil
import tempfile
import uuid
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from knowresearch.adapters import LocalObjectStorage
from knowresearch.adapters.memory import RedisSessionStore
from knowresearch.bootstrap import (
    build_production_document_store,
    build_production_ingest_pipeline,
    build_production_long_term_memory,
    build_production_memory_candidate_store,
    build_production_trace_store,
    build_production_vector_store,
    build_profile_extractor,
    build_profile_memory_service,
    build_profile_review_service,
    build_qa_engine,
)
from knowresearch.core.ports import (
    DocumentStorePort,
    EmbeddingProvider,
    FullTextStorePort,
    LLMProvider,
    RerankerProvider,
    VectorStorePort,
)
from knowresearch.core.schemas import Chunk, Document, Record, SectionBlock


pytestmark = pytest.mark.skipif(
    os.getenv("RUN_P2_INTEGRATION") != "1",
    reason="需要本地 PostgreSQL/Redis 服务，设置 RUN_P2_INTEGRATION=1 后运行",
)


REDIS_URL = os.getenv("P2_REDIS_URL", "redis://localhost:6379/11")


class _MockEmbedding(EmbeddingProvider):
    @property
    def dimension(self) -> int:
        return 3

    def embed_documents(self, texts):
        return [[0.1, 0.2, 0.3] for _ in texts]

    def embed_query(self, text):
        return [0.1, 0.2, 0.3]


class _MockVectorStore(VectorStorePort):
    def __init__(self, chunks):
        self._chunks = chunks

    def add(self, records, embeddings):
        pass

    def search(self, query_embedding, top_k=10):
        return [(chunk, 1.0) for chunk in self._chunks[:top_k]]

    def delete(self, record_ids):
        pass

    def clear(self):
        pass


class _MockFullTextStore(FullTextStorePort):
    def __init__(self, chunks):
        self._chunks = chunks

    def add(self, records):
        pass

    def search(self, query, top_k=10):
        return [(chunk, 1.0) for chunk in self._chunks[:top_k]]

    def delete(self, record_ids):
        pass

    def clear(self):
        pass


class _MockReranker(RerankerProvider):
    def rerank(self, query, documents, top_k=5):
        return [(index, 1.0) for index in range(min(top_k, len(documents)))]


class _MockLLM(LLMProvider):
    def generate(self, prompt, system_prompt="", temperature=0.7, max_tokens=2048):
        return "历史摘要"

    def generate_with_messages(self, messages, temperature=0.7, max_tokens=2048):
        return "这是答案[1]。"


class _ExtractionLLM(LLMProvider):
    def __init__(self, extraction_response: str):
        self.extraction_response = extraction_response
        self.calls: list[list[dict[str, str]]] = []

    def generate(self, prompt, system_prompt="", temperature=0.7, max_tokens=2048):
        return self.extraction_response

    def generate_with_messages(self, messages, temperature=0.7, max_tokens=2048):
        self.calls.append(messages)
        return "这是答案[1]。"


class _MockDocumentStore(DocumentStorePort):
    def __init__(self, doc, section):
        self._doc = doc
        self._section = section

    def save_document(self, doc):
        self._doc = doc

    def get_document(self, doc_id):
        return self._doc if self._doc.id == doc_id else None

    def get_document_by_hash(self, file_hash):
        return self._doc if self._doc.file_hash == file_hash else None

    def list_documents(self):
        return [self._doc]

    def save_sections(self, sections):
        pass

    def get_sections(self, doc_id):
        return [self._section] if self._section.document_id == doc_id else []

    def delete_document(self, doc_id):
        pass


def _make_chunk():
    return Chunk(
        id="chunk-p2",
        content="P2 短期记忆测试内容",
        metadata={"document_id": "doc-p2", "section_id": "sec-p2", "page": 1},
    )


def _make_engine(
    session_store,
    chunks,
    max_history_messages=20,
    trace_store=None,
    memory_context=None,
    llm_provider=None,
):
    doc = Document(id="doc-p2", metadata={"file_name": "p2.pdf"})
    section = SectionBlock(
        id="sec-p2",
        metadata={"document_id": "doc-p2", "section_title": "测试章节"},
    )
    engine = build_qa_engine(
        vector_store=_MockVectorStore(chunks),
        fulltext_store=_MockFullTextStore(chunks),
        document_store=_MockDocumentStore(doc, section),
        embedding_provider=_MockEmbedding(),
        reranker_provider=_MockReranker(),
        llm_provider=llm_provider or _MockLLM(),
        session_store=session_store,
        trace_store=trace_store,
        memory_context=memory_context,
    )
    engine.max_history_messages = max_history_messages
    return engine


def test_production_storage_builders_integration():
    doc_store = build_production_document_store()
    trace_store = build_production_trace_store()
    vector_store = build_production_vector_store(dimension=3)

    doc_id = f"doc-{uuid.uuid4().hex}"
    trace_id = f"trace-{uuid.uuid4().hex}"
    record_id = f"chunk-{uuid.uuid4().hex}"

    try:
        doc = Document(id=doc_id, metadata={"file_name": "p2-builder.pdf"})
        doc_store.save_document(doc)
        assert doc_store.get_document(doc_id).file_name == "p2-builder.pdf"

        trace_store.save(Record(id=trace_id, kind="trace", content="P2 trace"))
        assert trace_store.get(trace_id).content == "P2 trace"

        vector_store.add(
            [Record(id=record_id, kind="chunk", content="P2 vector")],
            [[0.1, 0.2, 0.3]],
        )
        assert any(record.id == record_id for record in vector_store.list_all())
    finally:
        doc_store.delete_document(doc_id)
        trace_store.delete(trace_id)
        vector_store.delete([record_id])


def test_redis_short_term_memory_with_qa_engine():
    session_store = RedisSessionStore(REDIS_URL)
    session_id = f"session-{uuid.uuid4().hex}"
    chunks = [_make_chunk()]
    engine = _make_engine(session_store, chunks, max_history_messages=2)

    try:
        engine.answer("第一轮问题", session_id=session_id)
        engine.answer("第二轮问题", session_id=session_id)

        messages = session_store.load_messages(session_id)
        assert len(messages) == 2
        assert messages[-2]["content"] == "第二轮问题"

        engine.answer("第三轮问题", session_id=session_id)
        assert session_store.load_summary(session_id) == "历史摘要"
    finally:
        session_store.delete_session(session_id)


def test_production_ingest_pipeline_writes_to_postgres():
    doc_store = build_production_document_store()
    vector_store = build_production_vector_store(dimension=3)
    temp_dir = Path(tempfile.mkdtemp(prefix="knowresearch_p2_ingest_"))
    docx_path = temp_dir / "p2_ingest.docx"
    file_store = LocalObjectStorage(str(temp_dir / "files"))

    docx = DocxDocument()
    docx.add_heading("P2 生产摄取测试", level=1)
    docx.add_paragraph("这是一段用于验证 PostgreSQL 生产摄取链路的内容。")
    docx.save(str(docx_path))

    pipeline = build_production_ingest_pipeline(
        document_store=doc_store,
        vector_store=vector_store,
        object_storage=file_store,
        embedding_provider=_MockEmbedding(),
    )
    result = pipeline.ingest_file(str(docx_path))
    chunk_ids: list[str] = []

    try:
        assert result.error == ""
        assert doc_store.get_document(result.document_id) is not None
        chunk_ids = [
            record.id
            for record in vector_store.list_all()
            if record.metadata.get("document_id") == result.document_id
        ]
        assert chunk_ids
    finally:
        doc_store.delete_document(result.document_id)
        vector_store.delete(chunk_ids)
        shutil.rmtree(temp_dir, ignore_errors=True)


def test_qa_engine_trace_writes_to_postgres():
    trace_store = build_production_trace_store()
    session_store = RedisSessionStore(REDIS_URL)
    session_id = f"session-{uuid.uuid4().hex}"
    engine = _make_engine(
        session_store=session_store,
        chunks=[_make_chunk()],
        trace_store=trace_store,
    )
    answer = engine.answer("P2-3 trace 测试问题", session_id=session_id, user_id="p2-user")

    try:
        trace = trace_store.get(answer.trace_id)
        assert trace is not None
        assert trace.kind == "trace"
        assert trace.source == session_id
        assert trace.metadata["status"] == "success"
        assert trace.metadata["user_id"] == "p2-user"
        assert trace.metadata["answer"] == "这是答案[1]。"
        assert trace.metadata["citation_count"] == 1
    finally:
        trace_store.delete(answer.trace_id)
        session_store.delete_session(session_id)


def test_postgres_long_term_profile_store_integration():
    store = build_production_long_term_memory()
    user_a = f"user-a-{uuid.uuid4().hex}"
    user_b = f"user-b-{uuid.uuid4().hex}"

    profile = store.upsert_profile(
        user_id=user_a,
        content="用户研究船舶水动力学",
        dimension="research_field",
        confidence=0.8,
    )

    try:
        got = store.get_profile(user_a, "research_field")
        assert got is not None
        assert got.content == "用户研究船舶水动力学"
        assert store.list(user_b, kind="profile") == []

        updated = store.upsert_profile(
            user_id=user_a,
            content="用户研究船舶水动力学和襟翼舵",
            dimension="research_field",
            confidence=0.9,
        )
        assert updated.id == profile.id
        assert "襟翼舵" in store.get_profile(user_a, "research_field").content
    finally:
        store.delete(profile.id, user_id=user_a)


def test_p2_memory_end_to_end_extract_review_inject():
    trace_store = build_production_trace_store()
    candidate_store = build_production_memory_candidate_store()
    long_term_memory = build_production_long_term_memory()
    session_store = RedisSessionStore(REDIS_URL)
    user_id = f"user-e2e-{uuid.uuid4().hex}"
    trace_id = f"trace-e2e-{uuid.uuid4().hex}"
    session_id = f"session-e2e-{uuid.uuid4().hex}"

    trace_store.save(
        Record(
            id=trace_id,
            kind="trace",
            content="我主要研究船舶水动力学，喜欢答案给出公式和页码。",
            metadata={
                "trace_id": trace_id,
                "user_id": user_id,
                "status": "success",
                "question": "我主要研究船舶水动力学，喜欢答案给出公式和页码。",
                "answer": "好的。",
            },
        )
    )

    llm = _ExtractionLLM(
        '[{"content":"用户研究船舶水动力学","dimension":"research_field",'
        '"confidence":0.9,"evidence":"用户自述研究方向"}]'
    )
    extractor = build_profile_extractor(
        llm_provider=llm,
        trace_store=trace_store,
        candidate_store=candidate_store,
    )
    candidates = extractor.extract_for_user(user_id)

    assert len(candidates) == 1
    candidate = candidates[0]
    assert candidate.status == "pending"
    assert candidate.source == trace_id

    review = build_profile_review_service(
        candidate_store=candidate_store,
        long_term_memory=long_term_memory,
    )
    profile = review.approve(candidate.id, user_id, review_note="确认")
    assert profile is not None

    memory_service = build_profile_memory_service(long_term_memory)
    context = memory_service.build_context(user_id, "优化方法")
    assert "船舶水动力学" in context

    engine = _make_engine(
        session_store=session_store,
        chunks=[_make_chunk()],
        trace_store=trace_store,
        memory_context=memory_service,
        llm_provider=llm,
    )
    answer = engine.answer("它有什么优化方法？", session_id=session_id, user_id=user_id)
    trace = trace_store.get(answer.trace_id)

    try:
        assert any(
            message["role"] == "system" and "船舶水动力学" in message["content"]
            for message in llm.calls[-1]
        )
        assert trace.metadata["memory_context_used"] is True
        assert "船舶水动力学" in trace.metadata["retrieval_query"]
    finally:
        trace_store.delete(trace_id)
        trace_store.delete(answer.trace_id)
        candidate_store.delete(candidate.id)
        long_term_memory.delete(profile.id, user_id=user_id)
        session_store.delete_session(session_id)
