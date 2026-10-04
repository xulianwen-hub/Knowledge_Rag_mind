"""Word 解析器：基于 python-docx 提取文本与章节结构。

Word 文档有原生标题样式（Heading 1/2/3），直接利用 style.name 识别层级。
"""

import hashlib
import os
import re

from docx import Document as DocxDocument

from knowresearch.core.schemas import Document, SectionBlock
from knowresearch.ingestion.parsers.base import ParserPort


REFERENCE_KEYWORDS = ["参考文献", "References", "REFERENCES", "Bibliography", "BIBLIOGRAPHY", "引用文献"]


def _compute_file_hash(file_path: str) -> str:
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _extract_heading_level(style_name: str) -> int | None:
    """从段落样式名提取标题层级。

    支持：
    - 英文：Heading 1, Heading 2, ...
    - 中文：标题 1, 标题 2, ...
    """
    m = re.search(r"(\d+)", style_name)
    if m and ("heading" in style_name.lower() or "标题" in style_name):
        return int(m.group(1))
    return None


class WordParser(ParserPort):
    """Word 解析器。"""

    def parse(self, file_path: str, doc_id: str = "") -> tuple[Document, list[SectionBlock]]:
        file_hash = _compute_file_hash(file_path)
        file_name = os.path.basename(file_path)

        docx = DocxDocument(file_path)

        raw_sections: list[dict] = []
        current: dict | None = None

        for para in docx.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            level = _extract_heading_level(para.style.name)
            if level is not None:
                if current is not None:
                    raw_sections.append(current)
                current = {
                    "level": level,
                    "title": text,
                    "content_lines": [],
                    "is_reference": any(kw in text for kw in REFERENCE_KEYWORDS),
                }
            else:
                if current is None:
                    current = {
                        "level": 1,
                        "title": "正文",
                        "content_lines": [],
                        "is_reference": False,
                    }
                current["content_lines"].append(text)

        if current is not None:
            raw_sections.append(current)

        doc_kwargs = dict(
            content="",
            source=file_path,
            metadata={
                "file_name": file_name,
                "file_hash": file_hash,
                "file_type": "docx",
                "total_pages": 0,
                "file_size": os.path.getsize(file_path),
            },
        )
        if doc_id:
            doc_kwargs["id"] = doc_id
        document = Document(**doc_kwargs)

        sections: list[SectionBlock] = []
        for idx, sec in enumerate(raw_sections):
            content = "\n".join(sec["content_lines"]).strip()
            section = SectionBlock(
                content=content,
                source=file_path,
                metadata={
                    "document_id": document.id,
                    "section_title": sec["title"],
                    "section_level": sec["level"],
                    "page_start": 0,
                    "page_end": 0,
                    "is_reference": sec["is_reference"],
                    "order_index": idx,
                },
            )
            sections.append(section)

        return document, sections