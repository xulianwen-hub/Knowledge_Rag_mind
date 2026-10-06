from __future__ import annotations

import os
from pathlib import Path

from docx import Document as DocxDocument

from knowresearch.core.schemas import MemoryCandidate, Record, SkillCandidate


ROOT = Path(__file__).resolve().parents[1]
RUNTIME_DIR = ROOT / "data" / "demo_runtime"
DOCS_DIR = RUNTIME_DIR / "docs"
FILES_DIR = RUNTIME_DIR / "files"


DEMO_DOCS = [
    {
        "file_name": "demo_flap_rudder.docx",
        "title": "襟翼舵优化演示",
        "paragraphs": [
            "襟翼舵通过主舵后缘的襟翼偏转来改变舵的水动力特性。",
            "襟翼参数会影响升力、阻力和失速特性，是船舶操纵性优化的重要方向。",
            "演示结论：比较襟翼舵方案时，应同时考虑升力增益、阻力代价和适用航速。",
        ],
    },
    {
        "file_name": "demo_pinn_fluid.docx",
        "title": "PINN 流体仿真演示",
        "paragraphs": [
            "物理信息神经网络把控制方程和边界条件加入损失函数。",
            "在流体力学中，PINN 可以用于正问题和逆问题求解。",
            "演示结论：PINN 适合在小样本、复杂边界或参数反演场景中与 CFD 互补。",
        ],
    },
    {
        "file_name": "demo_uuv_drag.docx",
        "title": "无人水下航行器阻力演示",
        "paragraphs": [
            "无人水下航行器的阻力优化通常关注艇型参数、外形曲线和推进效率。",
            "Myring 型回转体常用于外形优化和阻力对比研究。",
            "演示结论：阻力优化需要平衡总容积、平行中段长度和艏艉形状。",
        ],
    },
]


def main() -> None:
    os.environ["DEMO_MODE"] = "1"

    from knowresearch.adapters import (
        LocalObjectStorage,
    )
    from knowresearch.bootstrap import (
        build_embedding,
        build_ingest_pipeline,
        build_production_document_store,
        build_production_fulltext_store,
        build_production_memory_candidate_store,
        build_production_skill_candidate_store,
        build_production_skill_store,
        build_production_vector_store,
    )

    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    FILES_DIR.mkdir(parents=True, exist_ok=True)

    document_store = build_production_document_store()
    vector_store = build_production_vector_store()
    fulltext_store = build_production_fulltext_store(vector_store)
    object_storage = LocalObjectStorage(str(FILES_DIR))
    embedding = build_embedding()

    pipeline = build_ingest_pipeline(
        document_store=document_store,
        vector_store=vector_store,
        fulltext_store=fulltext_store,
        object_storage=object_storage,
        embedding_provider=embedding,
    )

    existing_names = {doc.file_name for doc in document_store.list_documents()}
    for doc_spec in DEMO_DOCS:
        if doc_spec["file_name"] in existing_names:
            print(f"skip existing: {doc_spec['file_name']}")
            continue
        path = DOCS_DIR / doc_spec["file_name"]
        _write_docx(path, doc_spec["title"], doc_spec["paragraphs"])
        result = pipeline.ingest_file(str(path))
        print(
            f"ingest: {doc_spec['file_name']} "
            f"chunks={result.chunk_count} error={result.error}"
        )

    _seed_memory_candidates(build_production_memory_candidate_store())
    _seed_skill_data(
        build_production_skill_candidate_store(),
        build_production_skill_store(),
    )
    print("\n演示数据准备完成。")
    print("下一步：")
    print('  $env:DEMO_MODE="1"')
    print("  python -m uvicorn knowresearch.api.app:app --reload")
    print("  python -m streamlit run frontend/app.py")


def _write_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    document = DocxDocument()
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(str(path))


def _seed_memory_candidates(store) -> None:
    profiles = [
        ("research_field", "用户主要研究船舶水动力学与襟翼舵优化"),
        ("output_preference", "用户希望答案给出公式和页码"),
    ]
    for index, (dimension, content) in enumerate(profiles, start=1):
        store.save(
            MemoryCandidate(
                id=f"demo-memory-{index}",
                user_id="demo-user",
                content=content,
                metadata={"dimension": dimension},
                source="demo",
                evidence="演示初始化数据",
                confidence=0.9,
                status="pending",
            )
        )


def _seed_skill_data(candidate_store, skill_store) -> None:
    candidate_store.save(
        SkillCandidate(
            id="demo-skill-candidate-1",
            user_id="demo-user",
            name="艇型对比分析",
            description="比较不同艇型优化方法的优缺点",
            trigger="用户要求比较不同艇型或优化方案",
            steps=["统一评价维度", "对比优缺点", "给出适用场景"],
            retrieval_params={"top_k": 10, "use_rerank": True},
            tags=["艇型", "对比"],
            source_trace_ids=["demo-trace-1", "demo-trace-2"],
            evidence="演示初始化数据",
            confidence=0.88,
            status="pending",
        )
    )
    skill_store.save(
        Record(
            id="demo-skill-1",
            kind="skill",
            content="PINN 方法对比分析",
            metadata={
                "user_id": "demo-user",
                "name": "PINN 方法对比分析",
                "description": "比较 PINN 在不同流体问题中的使用方式",
                "trigger": "用户要求比较 PINN 方法",
                "steps": ["识别任务类型", "比较损失函数", "比较数据和边界条件"],
                "retrieval_params": {"top_k": 10, "use_rerank": True},
                "tags": ["PINN", "对比"],
                "status": "approved",
                "version": 1,
                "usage_count": 3,
                "success_count": 2,
            },
        )
    )


if __name__ == "__main__":
    main()
